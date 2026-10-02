"""
07_build_slides.py
MFE 230GB Final Project - build the presentation draft slides/presentation.pptx

Charts are drawn from the files in output/ (scripts 02, 04, 05), so the deck
always matches the latest results. Speaker split is marked Person 1 / Person 2.

Requirements: pip install pandas matplotlib python-pptx
Run from the project root:  python code/07_build_slides.py
"""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT_A, OUT_C, OUT_R = Path("output/strategyA"), Path("output/strategyC"), Path("output/risk")
SLIDES = Path("slides")
IMG = SLIDES / "img"
IMG.mkdir(parents=True, exist_ok=True)

# palette: deep navy (dominant), amber accent, muted slate
NAVY, AMBER, SLATE = "14213D", "E39B17", "5C677D"
INK, WHITE, TINT = "1B1F27", "FFFFFF", "EEF1F6"
RED, GREEN = "B83227", "2E7D4F"
HEAD, BODY = "Cambria", "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------
def style(ax):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#B8BEC9")
    ax.tick_params(colors="#" + SLATE, labelsize=10)
    ax.grid(axis="y", color="#E4E7EC", lw=0.8)
    ax.set_axisbelow(True)


def save(fig, name):
    path = IMG / f"{name}.png"
    fig.tight_layout()
    fig.savefig(path, dpi=200, facecolor="white")
    plt.close(fig)
    return path


def cum(r):
    return (1 + r.fillna(0)).cumprod()


def make_charts():
    c = {}
    a = pd.read_csv(OUT_A / "monthly_returns.csv", index_col=0, parse_dates=True)
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    ax.plot(cum(a["Carry (net)"]), color="#" + NAVY, lw=1.8, label="Carry (net)")
    ax.plot(cum(a["Crowd-filtered (net)"]), color="#" + AMBER, lw=1.8, label="Crowd-filtered (net)")
    ax.axvline(pd.Timestamp("2011-01-01"), color="#" + SLATE, ls=":", lw=1)
    ax.text(pd.Timestamp("2011-03-01"), ax.get_ylim()[1] * 0.97, "out-of-sample", color="#" + SLATE, fontsize=9, va="top")
    ax.set_ylabel("Growth of $1", color="#" + SLATE)
    ax.legend(frameon=False, fontsize=10, loc="upper left")
    style(ax)
    c["a_cum"] = save(fig, "a_cum")

    # A vs LRV: rebuild cumulative of both over overlap from risk table inputs
    lrv_path = Path("data/raw/CurrencyPortfolios.xls")
    if lrv_path.exists():
        d = pd.read_excel(lrv_path, sheet_name="Developed currencies (net)", header=0)
        d.index = pd.to_datetime(d.iloc[:, 0]) + pd.offsets.MonthEnd(0)
        hml = d.filter(like="HML").iloc[:, 0]
        both = pd.concat([a["Carry (net)"], hml], axis=1, sort=True).dropna()
        fig, ax = plt.subplots(figsize=(7.4, 4.4))
        ax.plot(cum(both.iloc[:, 0]), color="#" + NAVY, lw=1.8, label="Our Strategy A carry (net)")
        ax.plot(cum(both.iloc[:, 1]), color="#" + SLATE, lw=1.8, ls="--", label="LRV developed HML (net)")
        ax.set_ylabel("Growth of $1", color="#" + SLATE)
        ax.legend(frameon=False, fontsize=10, loc="upper left")
        style(ax)
        c["a_lrv"] = save(fig, "a_lrv")

    crisis = a["Carry (net)"].loc["2008-07":"2009-03"] * 100
    fig, ax = plt.subplots(figsize=(5.6, 3.6))
    ax.bar(crisis.index.strftime("%b %y"), crisis.values,
           color=["#" + (RED if v < 0 else GREEN) for v in crisis.values])
    ax.set_ylabel("Monthly return (%)", color="#" + SLATE)
    ax.tick_params(axis="x", rotation=45)
    style(ax)
    c["a_2008"] = save(fig, "a_2008")

    cd = pd.read_csv(OUT_C / "daily_returns.csv", index_col=0, parse_dates=True)
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    s = cum(cd["Crypto carry (net)"])
    ax.plot(s, color="#" + NAVY, lw=1.5)
    ax.set_yscale("log")
    for d, lab in (("2020-03-12", "Covid"), ("2022-05-09", "LUNA"), ("2022-11-08", "FTX")):
        ax.axvline(pd.Timestamp(d), color="#" + RED, lw=1, ls=":")
        ax.text(pd.Timestamp(d), s.max() * 1.05, " " + lab, color="#" + RED, fontsize=10, va="bottom")
    ax.axvline(pd.Timestamp("2024-01-01"), color="#" + SLATE, ls=":", lw=1)
    ax.text(pd.Timestamp("2024-02-01"), s.min(), "out-of-sample", color="#" + SLATE, fontsize=9, va="bottom")
    ax.set_ylabel("Growth of $1 (net, log)", color="#" + SLATE)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FormatStrFormatter("%.1f"))
    ax.yaxis.set_minor_formatter(matplotlib.ticker.FormatStrFormatter("%.1f"))
    style(ax)
    c["c_cum"] = save(fig, "c_cum")

    fig, ax = plt.subplots(figsize=(6.4, 4.2))
    ax.plot(100 * cd["funding component"].cumsum(), color="#" + GREEN, lw=1.8, label="Funding received")
    ax.plot(100 * cd["price component"].cumsum(), color="#" + RED, lw=1.8, label="Price moves")
    ax.axhline(0, color="#B8BEC9", lw=0.8)
    ax.set_ylabel("Cumulative (% pts)", color="#" + SLATE)
    ax.legend(frameon=False, fontsize=10, loc="center right")
    style(ax)
    c["c_decomp"] = save(fig, "c_decomp")

    contrib = pd.read_csv(OUT_R / "C_contribution_by_coin.csv", index_col=0)["total P&L (% pts)"]
    fig, ax = plt.subplots(figsize=(6.0, 4.6))
    ax.barh(contrib.index, contrib.values,
            color=["#" + (RED if v < 0 else GREEN) for v in contrib.values])
    ax.axvline(0, color="#B8BEC9", lw=0.8)
    ax.set_xlabel("Total P&L contribution (% pts)", color="#" + SLATE)
    style(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color="#E4E7EC", lw=0.8)
    c["c_contrib"] = save(fig, "c_contrib")

    comb = pd.read_csv(OUT_R / "combined_monthly_returns.csv", index_col=0, parse_dates=True)
    fig, ax = plt.subplots(figsize=(8.0, 4.4))
    for col, color, lw in zip(comb.columns, ("#" + NAVY, "#" + SLATE, "#" + AMBER), (1.4, 1.4, 2.6)):
        ax.plot(cum(comb[col]), color=color, lw=lw, label=col)
    ax.set_ylabel("Growth of $1", color="#" + SLATE)
    ax.legend(frameon=False, fontsize=10, loc="upper left")
    style(ax)
    c["comb"] = save(fig, "comb")
    return c


# ---------------------------------------------------------------------------
# Slide helpers
# ---------------------------------------------------------------------------
W, H = 13.333, 7.5


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def text(slide, x, y, w, h, paras, size=16, color=INK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """paras: str or list of str / (str, dict) where dict may set size, color, bold, italic, bullet."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        txt, o = (p, {}) if isinstance(p, str) else p
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.space_after = Pt(o.get("space_after", space_after))
        run = par.add_run()
        run.text = ("•  " if o.get("bullet") else "") + txt
        f = run.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", False)
        f.color.rgb = rgb(o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill=TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def title(slide, t, sub=None, dark=False):
    text(slide, 0.6, 0.45, 10.6, 0.8, t, size=34, font=HEAD, bold=True,
         color=WHITE if dark else NAVY)
    if sub:
        text(slide, 0.6, 1.2, 11.5, 0.5, sub, size=16, color="C9D1E0" if dark else SLATE)


def speaker(slide, who, dark=False):
    s = box(slide, W - 2.1, 0.5, 1.5, 0.42, fill=AMBER)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = who
    r.font.size, r.font.bold, r.font.name = Pt(12), True, BODY
    r.font.color.rgb = rgb(NAVY)


def stat(slide, x, y, w, value, label, color=NAVY, fill=TINT, h=1.55):
    box(slide, x, y, w, h, fill=fill)
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.75, value, size=34, font=HEAD, bold=True, color=color)
    text(slide, x + 0.25, y + 0.92, w - 0.5, h - 1.0, label, size=12, color=SLATE)


def numbered(slide, x, y, w, items, gap=1.2):
    for i, (head, body) in enumerate(items):
        yy = y + i * gap
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy), Inches(0.55), Inches(0.55))
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(NAVY)
        c.line.fill.background()
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        r.font.size, r.font.bold, r.font.name = Pt(16), True, BODY
        r.font.color.rgb = rgb(WHITE)
        text(slide, x + 0.8, yy - 0.02, w - 0.8, 0.4, head, size=17, bold=True, color=NAVY)
        text(slide, x + 0.8, yy + 0.38, w - 0.8, gap - 0.4, body, size=14, color=INK)


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def table(slide, x, y, w, rows, col_w=None, size=12, header_fill=NAVY, row_h=0.36):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(r, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if r == 0 else (WHITE if r % 2 else TINT))
            cell.margin_left = cell.margin_right = Inches(0.08)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = str(val)
            run.font.size, run.font.name = Pt(size), BODY
            run.font.bold = r == 0
            run.font.color.rgb = rgb(WHITE if r == 0 else INK)
    return tbl


# ---------------------------------------------------------------------------
# Deck
# ---------------------------------------------------------------------------
def build():
    ch = make_charts()
    pa = pd.read_csv(OUT_A / "performance_table.csv", index_col=[0, 1])
    pc = pd.read_csv(OUT_C / "performance_table.csv", index_col=[0, 1])
    cv = pd.read_csv(OUT_A / "crowded_vs_not.csv", index_col=0)
    reg = pd.read_csv(OUT_R / "A_factor_regressions.csv", index_col=[0, 1])
    lrv = pd.read_csv(OUT_R / "A_vs_LRV_HML.csv", index_col=0)
    dec = pd.read_csv(OUT_C / "return_decomposition.csv", index_col=0)
    grid = pd.read_csv(OUT_C / "robustness_grid.csv")
    ctl = pd.read_csv(OUT_R / "C_risk_control.csv", index_col=[0, 1])
    cw = pd.read_csv(OUT_R / "C_worst_weeks.csv", index_col=0)
    comb = pd.read_csv(OUT_R / "combined_portfolio.csv", index_col=0)
    cm = pd.read_csv(OUT_R / "combined_monthly_returns.csv", index_col=0)
    a_crisis = pd.read_csv(OUT_A / "monthly_returns.csv", index_col=0, parse_dates=True)["Carry (net)"]
    crash08 = 100 * ((1 + a_crisis.loc["2008-07":"2009-03"]).prod() - 1)

    def S(df, k, col="Sharpe"):
        return df.loc[k, col]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]

    # 1 title -------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    text(s, 0.8, 2.0, 11.5, 1.2, "Crowded Carry", size=60, font=HEAD, bold=True, color=WHITE)
    text(s, 0.8, 3.25, 11.5, 0.8, "Crash risk in FX and crypto carry trades", size=26, color="C9D1E0")
    text(s, 0.8, 5.3, 11.5, 0.5, "Reza Zamani  ·  Paraj", size=18, color=AMBER, bold=True)
    text(s, 0.8, 5.85, 11.5, 0.5, "MFE 230GB Currency Markets  ·  Final project  ·  October 8, 2026",
         size=14, color="C9D1E0")
    notes(s, "Person 1 opens. One sentence: we test whether carry trades crash when they are crowded, in G10 FX and in crypto.")

    # 2 idea --------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Carry pays because it crashes", "Same economic mechanism, two very different markets")
    speaker(s, "Person 1")
    text(s, 0.6, 2.0, 5.4, 3.8, [
        ("High-yield assets earn a premium over low-yield ones.", {"bullet": True}),
        ("The premium is compensation for crash risk: when many investors hold the same trade, "
         "forced unwinding produces sudden losses (Brunnermeier, Nagel & Pedersen 2008).", {"bullet": True}),
        ("Hypothesis: crowding is measurable, and crowded carry should crash more.", {"bullet": True, "bold": True}),
    ], size=17, space_after=14)
    for i, (k, v1, v2) in enumerate([
            ("", "A: G10 FX", "C: Crypto"),
            ("Carry signal", "3-month rate differential", "Perpetual funding rate"),
            ("Crowding data", "CFTC speculator positions", "Funding = leveraged demand"),
            ("Universe", "7 currencies vs USD", "12 coins incl. LUNA, FTT"),
            ("Frequency", "Monthly", "Weekly rebalance, daily P&L")]):
        yy = 2.0 + i * 0.78
        if i == 0:
            text(s, 8.55, yy, 2.0, 0.5, v1, size=16, bold=True, color=NAVY)
            text(s, 10.65, yy, 2.1, 0.5, v2, size=16, bold=True, color=AMBER)
            continue
        box(s, 6.5, yy - 0.12, 6.25, 0.66)
        text(s, 6.7, yy, 1.8, 0.5, k, size=13, bold=True, color=SLATE)
        text(s, 8.55, yy, 2.0, 0.5, v1, size=13)
        text(s, 10.65, yy, 2.0, 0.5, v2, size=13)
    notes(s, "Explain why both are carry: an FX forward earns the rate differential; a short perpetual earns the funding rate.")

    # 3 strategy A rule ---------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Strategy A: crowding-filtered G10 carry", "JPY, EUR, GBP, CHF, CAD, AUD, NZD vs USD, monthly, 1999–2026")
    speaker(s, "Person 1")
    numbered(s, 0.6, 2.05, 7.2, [
        ("Carry portfolio", "Each month-end: long the 2 highest-rate currencies, short the 2 lowest. "
                            "Equal weights, dollar-neutral, 3 bp costs."),
        ("Crowding signal (alternative data)", "CFTC net speculative position / open interest, 36-month z-score, "
                                               "used only after the Friday release."),
        ("Portfolio crowding", "Average z of the long leg minus average z of the short leg."),
        ("Filter", "Crowding > 1: cut positions to half. All parameters fixed before testing; "
                   "out-of-sample from 2011."),
    ], gap=1.2)
    stat(s, 8.6, 2.05, 4.1, "1986", "CFTC positioning history (after fixing the pre-2000 CME exchange name)")
    stat(s, 8.6, 3.85, 4.1, "33", "crowded months out of 247 with a crowding signal")
    notes(s, "Mention the data issue we fixed: before Aug 2000 CME currency futures are reported as 'International Monetary Market'.")

    # 4 A results ---------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Crowding filter helps out of sample, not in 2008")
    speaker(s, "Person 1")
    picture(s, ch["a_cum"], 0.5, 1.45, w=7.9)
    stat(s, 8.75, 1.55, 4.0, f"{S(pa, ('Carry (net)', 'Out-of-sample')):.2f} → "
                             f"{S(pa, ('Crowd-filtered (net)', 'Out-of-sample')):.2f}",
         "Out-of-sample Sharpe, net: carry vs crowd-filtered (2011–2026)")
    stat(s, 8.75, 3.3, 4.0, f"{cv.loc['Not crowded', 'Next-month carry mean (%)']:.2f}% vs "
                            f"{cv.loc['Crowded', 'Next-month carry mean (%)']:.2f}%",
         "Next-month carry after non-crowded vs crowded months")
    stat(s, 8.75, 5.05, 4.0, f"{pa.loc[('Carry (net)', 'Full'), 'Max drawdown (%)']:.0f}%",
         "Max drawdown either way: the 2008 crash began from a non-crowded reading", color=RED)
    notes(s, f"Full-sample net Sharpe {S(pa, ('Carry (net)', 'Full')):.2f} vs {S(pa, ('Crowd-filtered (net)', 'Full')):.2f}. "
             "The filter roughly doubles turnover.")

    # 5 LRV check ---------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Our carry matches the academic carry factor",
          "Strategy A vs Lustig–Roussanov–Verdelhan developed HML, net, 1999–2021")
    speaker(s, "Person 1")
    if "a_lrv" in ch:
        picture(s, ch["a_lrv"], 0.5, 1.9, w=7.4)
    stat(s, 8.5, 1.95, 4.2, f"{lrv['Correlation'].iloc[0]:.2f}", "Monthly correlation with LRV developed HML")
    table(s, 8.5, 3.85, 4.2, [
        ["Net, 1999–2021", "Ours", "LRV"],
        ["Ann. mean (%)", f"{lrv.iloc[0]['Ann. mean (%)']:.1f}", f"{lrv.iloc[1]['Ann. mean (%)']:.1f}"],
        ["Ann. vol (%)", f"{lrv.iloc[0]['Ann. vol (%)']:.1f}", f"{lrv.iloc[1]['Ann. vol (%)']:.1f}"],
        ["Sharpe", f"{lrv.iloc[0]['Sharpe']:.2f}", f"{lrv.iloc[1]['Sharpe']:.2f}"],
        ["Max drawdown (%)", f"{lrv.iloc[0]['Max drawdown (%)']:.0f}", f"{lrv.iloc[1]['Max drawdown (%)']:.0f}"],
    ], col_w=[2.0, 1.1, 1.1], size=13)
    notes(s, "Course file CurrencyPortfolios.xls. This validates our data pipeline and code.")

    # 6 A risk ------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Strategy A loses when volatility spikes", "Monthly net carry regressed on risk factors, Newey-West t-stats")
    speaker(s, "Person 1")
    m = "RX (dollar) + dVol (LRV)"
    stat(s, 0.6, 1.95, 3.0, f"{reg.loc[(m, 'dVol (LRV)'), 'coef']:.2f}",
         f"Beta to equity-volatility changes (t = {reg.loc[(m, 'dVol (LRV)'), 'NW t-stat']:.1f})", color=RED)
    stat(s, 3.85, 1.95, 3.0, f"{reg.loc[(m, 'RX (dollar)'), 'coef']:.2f}",
         f"Beta to the dollar factor (t = {reg.loc[(m, 'RX (dollar)'), 'NW t-stat']:.1f})")
    stat(s, 0.6, 3.75, 3.0, f"{100 * reg.loc[(m, 'const'), 'coef']:.2f}%",
         f"Monthly alpha (t = {reg.loc[(m, 'const'), 'NW t-stat']:.1f}): not significant")
    stat(s, 3.85, 3.75, 3.0, f"{reg.loc[(m, 'R2'), 'coef']:.2f}", "R² of the two-factor model")
    text(s, 0.6, 5.65, 6.3, 1.2, "Worst months: Oct 2008, Aug 2007, May 2000, Jan 2009, and the Jan 2015 Swiss franc de-peg.",
         size=14, color=SLATE)
    text(s, 7.4, 1.95, 5.3, 0.4, f"The 2008 carry crash: {crash08:.1f}% (Jul 2008 – Mar 2009)",
         size=16, bold=True, color=NAVY)
    picture(s, ch["a_2008"], 7.3, 2.45, w=5.5)
    notes(s, "The textbook carry-crash exposure: funding currencies (JPY, CHF) rally when volatility jumps.")

    # 7 Strategy C rule ---------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Strategy C: crypto funding carry", "12 Binance USDT perpetuals, 2020–2026; USDC as a sanity check only")
    speaker(s, "Person 2")
    numbered(s, 0.6, 2.05, 7.2, [
        ("Signal", "Trailing 7-day average funding rate, data up to the day before."),
        ("Portfolio", "Each Sunday: short the top third, long the bottom third by funding. "
                      "Equal weights, dollar-neutral, 5 bp costs."),
        ("P&L", "Side × price return − side × funding paid. A short receives positive funding."),
        ("No survivorship bias", "LUNA and FTT included; a coin trades only on days with a price and volume > 0."),
    ], gap=1.2)
    stat(s, 8.6, 2.05, 4.1, "≈ 12–15%", "Average annual funding on BTC, ETH, XRP, LTC, LINK: the crypto 'interest rate'")
    stat(s, 8.6, 3.85, 4.1, "2022-11-14", "FTT's last traded day; the archive shows a frozen price afterwards, which we exclude")
    notes(s, "Funding is paid every 8 hours between longs and shorts and keeps the perpetual close to spot. "
             "High funding = crowded leveraged longs.")

    # 8 C results ---------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Crypto carry: weak in-sample, strong since 2024")
    speaker(s, "Person 2")
    picture(s, ch["c_cum"], 0.5, 1.45, w=7.9)
    stat(s, 8.75, 1.55, 4.0, f"{S(pc, ('Crypto carry (net)', 'Full')):.2f}",
         f"Full-sample net Sharpe ({pc.loc[('Crypto carry (net)', 'Full'), 'Ann. mean (%)']:.1f}% a year, "
         f"{pc.loc[('Crypto carry (net)', 'Full'), 'Ann. vol (%)']:.0f}% vol)")
    stat(s, 8.75, 3.3, 4.0, f"{S(pc, ('Crypto carry (net)', 'Out-of-sample')):.2f}",
         "Out-of-sample net Sharpe, 2024–2026")
    stat(s, 8.75, 5.05, 4.0, f"{cw['Base'].iloc[0]:.0f}%",
         "Worst week (LUNA, May 2022): the strategy was long LUNA", color=RED)
    notes(s, "LUNA's funding had turned negative as traders shorted it, so the rule held it long into the collapse. "
             "FTX week: -11.6%.")

    # 9 C decomposition + robustness -------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Funding is earned, but price moves dominate", "Return decomposition and robustness")
    speaker(s, "Person 2")
    picture(s, ch["c_decomp"], 0.5, 1.85, w=6.1)
    rows = [["% per year", "Funding", "Price", "Net"]]
    for p in ("Full", "In-sample", "Out-of-sample"):
        rows.append([p, f"{dec.loc[p, 'Funding carry (% p.a.)']:+.1f}",
                     f"{dec.loc[p, 'Price component (% p.a.)']:+.1f}", f"{dec.loc[p, 'Net total (% p.a.)']:+.1f}"])
    table(s, 7.0, 1.95, 5.8, rows, col_w=[2.2, 1.2, 1.2, 1.2], size=13)
    rows = [["Lookback × rebalance", "Sharpe net", "Sharpe OOS"]]
    for _, r in grid.iterrows():
        tag = " (base)" if r["lookback_days"] == 7 and r["rebalance"] == "weekly" else ""
        rows.append([f"{r['lookback_days']} days, {r['rebalance']}{tag}", f"{r['Sharpe net']:.2f}",
                     f"{r['Sharpe OOS net']:.2f}"])
    table(s, 7.0, 3.75, 5.8, rows, col_w=[3.0, 1.4, 1.4], size=12, row_h=0.33)
    notes(s, "Shorter signals do better even after costs: funding information decays within days. "
             "We report the pre-specified 7-day weekly case as the main result and do not switch after seeing results.")

    # 10 C risk control ---------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "A pre-specified risk control did not help", "Cap each coin at 1/6 of a leg; no longs with funding below −50% a year")
    speaker(s, "Person 2")
    picture(s, ch["c_contrib"], 0.5, 1.85, w=5.6)
    table(s, 6.5, 1.95, 6.3, [
        ["Net", "Base", "Controlled"],
        ["Sharpe, full", f"{ctl.loc[('Base', 'Full'), 'Sharpe']:.2f}", f"{ctl.loc[('Risk-controlled', 'Full'), 'Sharpe']:.2f}"],
        ["Sharpe, out-of-sample", f"{ctl.loc[('Base', 'Out-of-sample'), 'Sharpe']:.2f}",
         f"{ctl.loc[('Risk-controlled', 'Out-of-sample'), 'Sharpe']:.2f}"],
        ["Max drawdown (%)", f"{ctl.loc[('Base', 'Full'), 'Max drawdown (%)']:.0f}",
         f"{ctl.loc[('Risk-controlled', 'Full'), 'Max drawdown (%)']:.0f}"],
        ["LUNA week (%)", f"{cw['Base'].iloc[0]:.1f}", f"{cw['Risk-controlled'].iloc[0]:.1f}"],
    ], col_w=[2.9, 1.7, 1.7], size=13)
    text(s, 6.5, 4.05, 6.3, 2.8, [
        ("The cap mostly halves exposure (3–4 coins per leg).", {"bullet": True}),
        ("LUNA's funding crossed −50% mid-week, after the Sunday rebalance.", {"bullet": True}),
        ("The filter blocked the biggest winners, BNB and SOL.", {"bullet": True}),
        ("Lesson: deeply negative funding means crowded shorts, not distress.", {"bullet": True, "bold": True, "color": NAVY}),
    ], size=15, space_after=8)
    notes(s, "Losses are concentrated: LUNA -87, XRP -58, DOGE -44 percentage points. Beta to BTC is about zero (t = -0.75).")

    # 11 combined ---------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Combining FX and crypto carry", "Each scaled to 10% vol with lagged volatility, 50/50, 2021–2026")
    speaker(s, "Person 2")
    picture(s, ch["comb"], 0.5, 1.85, w=7.8)
    corr = cm.iloc[:, 0].corr(cm.iloc[:, 1])
    stat(s, 8.75, 1.95, 4.0, f"{corr:.2f}", "Correlation between A and C")
    table(s, 8.75, 3.75, 4.0, [
        ["", "Sharpe", "Vol (%)", "MaxDD (%)"],
        *[[k.split(" (")[0], f"{comb.loc[k, 'Sharpe']:.2f}", f"{comb.loc[k, 'Ann. vol (%)']:.1f}",
           f"{comb.loc[k, 'Max drawdown (%)']:.0f}"] for k in comb.index],
    ], col_w=[1.45, 0.8, 0.85, 0.9], size=12)
    notes(s, "The two carry trades crash at different times, so the combination is much smoother; "
             "Strategy C was weak in this window, so the combination does not beat A on Sharpe.")

    # 12 conclusions ------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "Conclusions and limitations", dark=True)
    speaker(s, "Both")
    text(s, 0.6, 1.7, 5.9, 0.5, "What we found", size=20, bold=True, color=AMBER)
    text(s, 0.6, 2.3, 5.9, 4.5, [
        ("FX carry reproduces the academic factor (corr 0.82) and is exposed to volatility shocks.", {"bullet": True}),
        ("Crowding predicts weaker carry and modestly improves the out-of-sample Sharpe.", {"bullet": True}),
        ("Crypto funding carry is market-neutral, but its crashes are coin-specific: it ends up long what traders short hardest.", {"bullet": True}),
        ("FX and crypto carry are nearly uncorrelated (0.11).", {"bullet": True}),
    ], size=15, color=WHITE, space_after=12)
    text(s, 7.0, 1.7, 5.7, 0.5, "Limitations", size=20, bold=True, color=AMBER)
    text(s, 7.0, 2.3, 5.7, 4.5, [
        ("Short crypto history (2020–2026) and only 12 coins.", {"bullet": True}),
        ("Few crowded FX months (33); the 2008 crash was not flagged.", {"bullet": True}),
        ("Crypto results depend on lookback and rebalancing frequency.", {"bullet": True}),
        ("Costs are simple estimates; no funding-rate caps, borrow limits or exchange risk.", {"bullet": True}),
    ], size=15, color=WHITE, space_after=12)
    notes(s, "Close with the main message: carry is compensation for crash risk, and crowding is measurable in both markets.")

    path = SLIDES / "presentation.pptx"
    prs.save(path)
    print(f"Wrote {path} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
